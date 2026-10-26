"""A small, realistic construction library for the demo, the end-to-end tests and the evaluation set.

It covers the hard cases on purpose: an
A1 drawing with a title block, a scanned page (OCR), an RFI and its near-duplicate copy, Word minutes with an actions
table, Arabic minutes, an Excel bill of quantities, a PowerPoint progress report, an e-mail with a PDF attachment, the
same spec section in a second project with a different answer, and a restricted commercial document.

    rag demo seed   # writes it to ./data/demo and registers it as two sources
"""

from __future__ import annotations

import email.message
from pathlib import Path

PUBLIC = "public"
RESTRICTED = "restricted"


def _pdf(path: Path, pages: list[list[str]], *, pagesize: tuple[float, float] | None = None,
         title: str | None = None) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    size = pagesize or A4
    c = canvas.Canvas(str(path), pagesize=size)
    if title:
        c.setTitle(title)
    for lines in pages:
        y = size[1] - 60
        for line in lines:
            if line.startswith("## "):
                c.setFont("Helvetica-Bold", 12)
                line = line[3:]  # noqa: PLW2901
            else:
                c.setFont("Helvetica", 10)
            for part in _wrap(line, 95):
                c.drawString(56, y, part)
                y -= 15
            y -= 4
        c.showPage()
    c.save()


def _wrap(text: str, width: int) -> list[str]:
    words, lines, current = text.split(), [], ""
    for word in words:
        if len(current) + len(word) + 1 > width:
            lines.append(current)
            current = word
        else:
            current = f"{current} {word}".strip()
    if current or not lines:
        lines.append(current)
    return lines


SPEC_REV2 = [
    [
        "## SECTION 03 30 00",
        "## CAST-IN-PLACE CONCRETE",
        "Project: Marina Tower    Date: 10/02/2025",
        "## PART 1 - GENERAL",
        "## 1.01 SUMMARY",
        "A. This Section specifies cast-in-place concrete for columns, core walls, slabs and footings.",
        "B. Related Sections: Section 03 20 00 Concrete Reinforcing; Section 03 11 00 Concrete Forming.",
        "## 1.02 SUBMITTALS",
        "A. Submit mix designs for each concrete class at least 21 days before the first pour.",
        "## PART 2 - PRODUCTS",
        "## 2.01 CONCRETE MIXES",
        "A. Columns on Level 3 and above shall use concrete grade C40 with a maximum aggregate size of 20 mm.",
        "B. Slabs shall use concrete grade C35. Footings shall use concrete grade C40 with sulphate resisting cement.",
        "C. Maximum water-cement ratio for grade C40 concrete is 0.40.",
    ],
    [
        "## PART 3 - EXECUTION",
        "## 3.01 PLACING",
        "A. Do not place concrete when the ambient temperature exceeds 35 degrees C unless the Engineer approves "
        "hot-weather measures in writing.",
        "B. Concrete cover to reinforcement shall be 50 mm for columns and 25 mm for internal slabs.",
        "## 3.02 CURING",
        "A. Keep exposed surfaces continuously moist for at least 7 days after placing.",
        "B. Remove column formwork no earlier than 24 hours after placing.",
    ],
]

HARBOUR_SPEC = [
    [
        "## SECTION 03 30 00",
        "## CAST-IN-PLACE CONCRETE",
        "Project: Harbour View    Date: 01/03/2025",
        "## PART 2 - PRODUCTS",
        "## 2.01 CONCRETE MIXES",
        "A. Columns at all levels shall use concrete grade C50 with silica fume.",
        "B. Marine-exposed elements shall use concrete grade C50 with corrosion inhibitor.",
    ],
]

RFI_023 = [[
    "## REQUEST FOR INFORMATION No. 023",
    "Project: Marina Tower    Date: 03/03/2025    From: Gulf Build Contracting    To: Engineer",
    "Subject: Fire damper access panels on Level 3 corridors",
    "Question: The ceiling layout on drawing A-301 shows no access to the fire dampers above the Level 3 corridors. "
    "Please confirm the access arrangement.",
    "Response: Provide 600 x 600 mm access panels at each fire damper on the Level 3 corridors, located within "
    "500 mm of the damper. Panels shall have the same fire rating as the ceiling.",
    "Responded by: M. Haddad, Mechanical Engineer, 06/03/2025",
]]

RFI_023_COPY = [[
    "## REQUEST FOR INFORMATION No. 023",
    "Date: 03/03/2025    From: Gulf Build Contracting    To: Engineer",
    "Subject: Fire damper access panels on Level 3 corridors",
    "Question: The ceiling layout on drawing A-301 shows no access to the fire dampers above the Level 3 corridors. "
    "Please confirm the access arrangement.",
    "Response: Provide 600 x 600 mm access panels at each fire damper on the Level 3 corridors, located within "
    "500 mm of the damper. Panels shall have the same fire rating as the ceiling.",
    "Responded by: M. Haddad, Mechanical Engineer, 06/03/2025 (filed copy)",
]]

VO_15 = [[
    "## VARIATION ORDER VO-15",
    "Project: Marina Tower    Date: 18/03/2025",
    "Description: Additional bored piles at grid lines F-H due to revised soil report.",
    "Agreed value of variation: AED 1,450,000 (one million four hundred fifty thousand dirhams).",
    "Time extension: 21 calendar days.",
]]

SUB_084_ATTACHMENT = [[
    "## SUBMITTAL SUB-084 - REBAR SHOP DRAWINGS",
    "Project: Marina Tower",
    "Shop drawings for Level 3 column reinforcement, columns C1 to C24.",
    "Main bars: 12 bars of 25 mm diameter per column; links 10 mm at 150 mm centres.",
]]


def _drawing(path: Path) -> None:
    from reportlab.lib.pagesizes import A1, landscape
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = landscape(A1)
    c = canvas.Canvas(str(path), pagesize=(width, height))
    c.setTitle("Level 03 Column Layout")
    c.setLineWidth(0.5)
    # Grid lines and columns: enough vector paths to read as a drawing sheet.
    for i in range(24):
        x = 120 + i * 70
        c.line(x, 300, x, height - 120)
        c.drawString(x - 4, height - 110, chr(65 + i % 26))
    for j in range(16):
        y = 300 + j * 75
        c.line(100, y, width - 700, y)
        c.drawString(90, y - 3, str(j + 1))
    for i in range(0, 24, 2):
        for j in range(0, 16, 3):
            c.rect(120 + i * 70 - 8, 300 + j * 75 - 8, 16, 16, fill=1)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(120, 240, "GENERAL NOTES")
    c.setFont("Helvetica", 11)
    notes = ["1. ALL COLUMNS C1 TO C24 ON LEVEL 03 ARE 600 x 600 mm.",
             "2. CONCRETE GRADE C40 FOR ALL COLUMNS. REFER TO SPECIFICATION SECTION 03 30 00.",
             "3. COVER TO REINFORCEMENT 50 mm. REFER TO SUBMITTAL SUB-084 FOR REBAR DETAILS."]
    for k, note in enumerate(notes):
        c.drawString(120, 220 - k * 16, note)
    # Title block, bottom-right.
    x0, y0 = width - 640, 30
    c.rect(x0, y0, 610, 240)
    labels = [("PROJECT:", "MARINA TOWER"), ("DRAWING TITLE:", "LEVEL 03 COLUMN LAYOUT"),
              ("DRAWING NO:", "PRJ1-ABC-ZZ-03-DR-S-2001"), ("REV:", "C02"), ("SCALE:", "1:100"),
              ("DATE:", "14/02/2025"), ("STATUS:", "IFC")]
    for k, (label, value) in enumerate(labels):
        y = y0 + 210 - k * 28
        c.setFont("Helvetica-Bold", 11)
        c.drawString(x0 + 14, y, label)
        c.setFont("Helvetica", 11)
        c.drawString(x0 + 170, y, value)
    c.showPage()
    c.save()


def _scanned_pdf(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    font_path = next((p for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                                  "/usr/share/fonts/dejavu/DejaVuSans.ttf") if Path(p).is_file()), None)
    big = ImageFont.truetype(font_path, 34) if font_path else ImageFont.load_default()
    body = ImageFont.truetype(font_path, 26) if font_path else ImageFont.load_default()
    draw.text((90, 110), "Site Instruction No. 14", fill="black", font=big)
    lines = ["Project: Marina Tower", "Date: 12 March 2025",
             "Install fire dampers at all shaft penetrations on Level 3.",
             "Seal gaps around ducts with fire rated sealant.", "Issued by: Resident Engineer"]
    for k, line in enumerate(lines):
        draw.text((90, 220 + k * 60), line, fill="black", font=body)
    image.save(path, "PDF", resolution=150.0)


def _docx_minutes(path: Path) -> None:
    import docx

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    doc.core_properties.title = "Minutes of Meeting 14"
    doc.add_heading("Minutes of Meeting No. 14 - Marina Tower", level=1)
    doc.add_paragraph("Date: 14/03/2025. Location: site office. Attendees: client, engineer, contractor.")
    doc.add_heading("Discussion", level=2)
    doc.add_paragraph("The contractor reported that Level 3 column pours are on schedule. The engineer reminded "
                      "the contractor that formwork must not be struck before 24 hours.")
    doc.add_heading("Actions", level=2)
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Action", "Owner", "Due"], strict=True):
        cell.text = text
    for row in (["Inspect Level 3 column formwork before each pour", "Ahmed Saleh (Contractor QA)", "20/03/2025"],
                ["Issue revised fire damper schedule", "M. Haddad (Engineer)", "25/03/2025"]):
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    doc.save(str(path))


def _docx_arabic(path: Path) -> None:
    import docx

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    doc.add_heading("محضر اجتماع رقم 15 - مشروع برج المارينا", level=1)
    doc.add_paragraph("التاريخ: 21/03/2025")
    doc.add_paragraph("تم الاتفاق على أن يكون موعد صب خرسانة سقف الطابق الثالث يوم 15 أبريل 2025 بعد اعتماد "
                      "المهندس لحديد التسليح.")
    doc.add_paragraph("طلب الاستشاري تقديم نتائج اختبار المكعبات للخرسانة من الفئة C40 قبل نهاية الشهر.")
    doc.save(str(path))


def _xlsx_boq(path: Path) -> None:
    import openpyxl

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Structural Works"
    ws.append(["Item", "Description", "Unit", "Qty", "Rate (AED)", "Amount (AED)"])
    rows = [["B-10", "Concrete grade C40 to columns", "m3", 1250, 520, 650000],
            ["B-11", "Concrete grade C35 to slabs", "m3", 3400, 480, 1632000],
            ["B-12", "Reinforcement steel 16mm high yield", "t", 420, 2850, 1197000],
            ["B-13", "Formwork to columns", "m2", 5200, 95, 494000]]
    for row in rows:
        ws.append(row)
    wb.properties.title = "Bill of Quantities - Structural Works"
    wb.save(str(path))


def _pptx_progress(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Marina Tower - Progress Report March 2025"
    body = slide.placeholders[1].text_frame
    body.text = "Overall progress: 42% complete against 45% planned."
    body.add_paragraph().text = "Structure: Level 3 columns poured; Level 3 slab scheduled for April."
    slide2 = deck.slides.add_slide(deck.slide_layouts[5])
    slide2.shapes.title.text = "Key risks"
    box = slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2)).text_frame
    box.text = "Late approval of SUB-084 rebar shop drawings could delay the Level 3 slab pour."
    deck.save(str(path))


def _email_with_attachment(path: Path, attachment_pdf: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    msg = email.message.EmailMessage()
    msg["From"] = "engineer@consultant.example"
    msg["To"] = "pm@contractor.example"
    msg["Date"] = "Mon, 10 Mar 2025 09:30:00 +0400"
    msg["Subject"] = "SUB-084 Rebar shop drawings - approved with comments"
    msg.set_content("Dear PM,\n\nSubmittal SUB-084 (Level 3 column rebar shop drawings) is approved with comments: "
                    "increase lap length to 50 diameters at column splices. Please resubmit the revised sheets for "
                    "record.\n\nRegards,\nEngineer")
    msg.add_attachment(attachment_pdf.read_bytes(), maintype="application", subtype="pdf",
                       filename="SUB-084 Rebar shop drawings.pdf")
    path.write_bytes(bytes(msg))


def build_demo_corpus(root: Path) -> dict[str, Path]:
    """Write the corpus under root/public and root/restricted. Returns the two source folders."""
    public = root / PUBLIC
    restricted = root / RESTRICTED
    mt = public / "Marina Tower"
    _pdf(mt / "Specifications" / "03 30 00 Cast-in-Place Concrete Rev 2.pdf", SPEC_REV2,
         title="Section 03 30 00 Cast-in-Place Concrete")
    _drawing(mt / "Drawings" / "Structural" / "PRJ1-ABC-ZZ-03-DR-S-2001_S4_C02.pdf")
    _pdf(mt / "RFIs" / "RFI-023 Fire damper access panels.pdf", RFI_023)
    _pdf(public / "Archive" / "RFI-023 Fire damper access panels (filed copy).pdf", RFI_023_COPY)
    _scanned_pdf(mt / "Site Instructions" / "SI-014 Site Instruction (scan).pdf")
    _docx_minutes(mt / "Minutes" / "Minutes of Meeting 14.docx")
    _docx_arabic(mt / "Minutes" / "محضر اجتماع 15.docx")
    _xlsx_boq(mt / "Commercial" / "BOQ Structural Works.xlsx")
    _pptx_progress(mt / "Reports" / "Progress Report March 2025.pptx")
    attachment = root / ".tmp" / "SUB-084.pdf"
    _pdf(attachment, SUB_084_ATTACHMENT)
    _email_with_attachment(mt / "Submittals" / "SUB-084 Rebar shop drawings approval.eml", attachment)
    _pdf(public / "Harbour View" / "Specifications" / "03 30 00 Cast-in-Place Concrete.pdf", HARBOUR_SPEC)
    _pdf(restricted / "Marina Tower" / "Commercial" / "VO-15 Variation Order - Additional piling.pdf", VO_15)
    return {"public": public, "restricted": restricted}

