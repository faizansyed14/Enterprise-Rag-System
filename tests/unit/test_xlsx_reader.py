from __future__ import annotations

import io

import openpyxl

from rag.ingestion import formats
from rag.ingestion.formats import read_xlsx


def _workbook(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = openpyxl.Workbook(write_only=True)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


HEADER = ["Ref", "Description", "Qty", "Status"]


def _rows(n: int) -> list[list[object]]:
    return [[f"ITM-{i:04d}", f"Item {i}", i % 9 + 1, "Ordered"] for i in range(n)]


def test_a_small_sheet_is_one_page_with_the_plain_label() -> None:
    doc = read_xlsx(_workbook({"Rates": [HEADER, *_rows(5)]}))
    assert len(doc.pages) == 1 and doc.pages[0].label == "Sheet: Rates"
    assert doc.pages[0].tables[0].rows[0] == HEADER and len(doc.pages[0].tables[0].rows) == 6


def test_a_long_sheet_is_cut_into_pages_that_each_repeat_the_header_and_name_their_rows(monkeypatch) -> None:
    monkeypatch.setattr(formats, "ROWS_PER_PAGE", 100)
    doc = read_xlsx(_workbook({"Register": [HEADER, *_rows(250)]}))
    assert [p.label for p in doc.pages] == ["Sheet: Register, rows 2-101", "Sheet: Register, rows 102-201", "Sheet: Register, rows 202-251"]
    assert all(p.tables[0].rows[0] == HEADER for p in doc.pages)
    assert sum(len(p.tables[0].rows) - 1 for p in doc.pages) == 250 and not doc.warnings  # nothing is dropped


def test_title_lines_above_the_header_are_text_not_column_names() -> None:
    doc = read_xlsx(_workbook({"Report": [["Project register"], ["Issued 2026-10-01"], [], HEADER, *_rows(3)]}))
    page = doc.pages[0]
    assert page.tables[0].rows[0] == HEADER and len(page.tables[0].rows) == 4
    assert "Project register" in page.text and "Issued 2026-10-01" in page.text


def test_every_sheet_of_a_workbook_is_read_and_empty_sheets_are_skipped() -> None:
    doc = read_xlsx(_workbook({"A": [HEADER, *_rows(2)], "Empty": [], "B": [HEADER, *_rows(2)]}))
    assert [p.label for p in doc.pages] == ["Sheet: A", "Sheet: B"] and [p.number for p in doc.pages] == [1, 2]


def test_a_sheet_of_figures_with_no_header_keeps_every_row_as_data_and_invents_no_column_names() -> None:
    doc = read_xlsx(_workbook({"Numbers": [[1, 2, 3], [4, 5, 6], [7, 8, 9]]}))
    table = doc.pages[0].tables[0]
    assert table.header_rows == 0 and table.rows == [["1", "2", "3"], ["4", "5", "6"], ["7", "8", "9"]]


def test_dated_sales_figures_do_not_turn_their_first_row_into_column_names() -> None:
    import datetime

    doc = read_xlsx(_workbook({"Sales": [[datetime.date(2023, 12, 1), 80237], [datetime.date(2024, 1, 1), 86369]]}))
    table = doc.pages[0].tables[0]
    assert table.header_rows == 0 and table.rows[0] == ["2023-12-01", "80237"] and len(table.rows) == 2


def test_a_sparse_first_row_of_labels_is_still_the_header() -> None:
    doc = read_xlsx(_workbook({"Prices": [["", "", "pvc", "fedrigoni", "", "", "quantity"], ["15 chocolates", "", 14, 17.4, "", "", 500]]}))
    assert doc.pages[0].tables[0].header_rows == 1 and doc.pages[0].tables[0].rows[0][2] == "pvc"


def test_a_form_row_of_labels_and_values_is_not_taken_for_a_header() -> None:
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Required information"
    ws.append(["Registrant Name", "", "Acme Trading LLC", "", "TRN number", 100123456700003, "Tax Return Period", "From : 01-07-2022 To 30-09-2022"])
    ws.append(["Requirements"])
    ws.append(["Please ensure you fill all the sheets"])
    buffer = io.BytesIO()
    wb.save(buffer)
    page = read_xlsx(buffer.getvalue()).pages[0]
    rows = page.tables[0].rows
    assert page.tables[0].header_rows == 0
    assert any("100123456700003" in cell for row in rows for cell in row) and any("Acme Trading LLC" in cell for cell in rows[0])
